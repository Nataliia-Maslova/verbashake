"""
scripts/generate_phrasebook_language.py -- one-time (resumable) translation
of the Phrasebook module's phrase database (data/vocabulary_translated.xlsx)
into a brand-new target language column.

Mirrors scripts/generate_grammar_language.py's approach (batch by lesson_id,
one Gemini call per lesson, resumable, backed-up workbook) but for
vocabulary_translated.xlsx's 29 sheets. Only the 25 non-Word-Bank sheets are
touched -- engine.vocab_loader.WORD_BANK_SHEETS (Basic/Verbs/Food/City) are
Vocabulary's old data source, dead for any target_lang since the 2026-08-21
CEFR-J switchover (grammar.py::_load_vocabulary loads CEFR-J now, not this
file) -- translating them would be wasted API cost for content nothing reads.

There is no topic/lesson_name column to translate per language here (unlike
Grammar's topic_en/topic_ro) -- "lesson_name" is a single shared English
index label, and the student-facing sheet name is translated lazily on the
fly via engine.gemini.translate_phrase() in path_app.py -- not this file.

Usage:
    python scripts/generate_phrasebook_language.py --lang Romanian --limit 3
    python scripts/generate_phrasebook_language.py --lang Romanian
        Resumable across ALL 25 sheets; --overwrite forces regeneration.

Requires GEMINI_API_KEY. Backs up the workbook once per day before writing
(data/vocabulary_translated.xlsx.bak_<YYYYMMDD>).
"""
from __future__ import annotations

import argparse
import datetime
import shutil
import sys
import time
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from engine import gemini
from engine.loader import LANG_COLUMNS
from engine.vocab_loader import WORD_BANK_SHEETS

DB_PATH = ROOT / "data" / "vocabulary_translated.xlsx"


def _backup_once() -> None:
    stamp = datetime.date.today().strftime("%Y%m%d")
    bak = DB_PATH.with_name(f"{DB_PATH.name}.bak_{stamp}")
    if not bak.exists():
        shutil.copy2(DB_PATH, bak)
        print(f"Backed up {DB_PATH.name} -> {bak.name}")


def _ensure_column(ws, headers: list, name: str) -> int:
    if name in headers:
        return headers.index(name) + 1
    col = len(headers) + 1
    ws.cell(row=1, column=col, value=name)
    headers.append(name)
    return col


def translate_lesson(phrases: list[str], lang: str) -> list[str]:
    numbered = "\n".join(f"{i+1}. {p}" for i, p in enumerate(phrases))
    prompt = (
        f"Translate these English phrases into {lang}, for a language-"
        f"learning app's phrasebook (everyday conversational phrases). "
        f"Keep the same register and any names (e.g. \"Alex\") as-is. "
        f"Return ONLY a JSON object, no markdown fences, no commentary:\n"
        f'{{"phrases": ["<translation 1>", "<translation 2>", ...]}}\n\n'
        f"Phrases:\n{numbered}"
    )
    result = gemini._model(gemini._LITE).generate_content(prompt)
    parsed = gemini._parse_json(gemini._safe_text(result), fallback=None)
    if not parsed or "phrases" not in parsed or len(parsed["phrases"]) != len(phrases):
        raise ValueError(f"bad/mismatched response: {parsed!r}")
    return parsed["phrases"]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--lang", required=True, help="Full language name, e.g. 'Romanian'.")
    parser.add_argument("--limit", type=int, default=None, help="Max lessons to translate this run.")
    parser.add_argument("--overwrite", action="store_true", help="Regenerate even already-filled lessons.")
    parser.add_argument("--sleep", type=float, default=0.0, help="Seconds to sleep between calls.")
    args = parser.parse_args()

    if args.lang not in LANG_COLUMNS:
        sys.exit(f"'{args.lang}' not in engine.loader.LANG_COLUMNS -- add it there first.")
    code = LANG_COLUMNS[args.lang]

    _backup_once()

    wb = load_workbook(DB_PATH)
    sheet_names = [s for s in wb.sheetnames if s not in WORD_BANK_SHEETS]

    done = 0
    for sheet in sheet_names:
        ws = wb[sheet]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        en_col = headers.index("en") + 1
        lesson_id_col = headers.index("lesson_id") + 1
        lang_col = _ensure_column(ws, headers, code)

        lessons: dict[int, list[int]] = {}
        for row in range(2, ws.max_row + 1):
            lid = ws.cell(row=row, column=lesson_id_col).value
            if lid is None:
                continue
            lessons.setdefault(int(lid), []).append(row)

        for lid in sorted(lessons):
            # A handful of rows across this workbook are fully empty (known
            # data-quality gap, LINGUISTIC_AUDIT.md) -- skip them here rather
            # than sending an empty string to Gemini, which just drops it
            # from its response and desyncs the zip() below.
            rows = [r for r in lessons[lid] if (ws.cell(row=r, column=en_col).value or "").strip()]
            if not rows:
                continue
            already_filled = all(
                (ws.cell(row=r, column=lang_col).value or "").strip() not in ("", "TODO")
                for r in rows
            )
            if already_filled and not args.overwrite:
                continue

            en_phrases = [ws.cell(row=r, column=en_col).value for r in rows]

            try:
                translated = translate_lesson(en_phrases, args.lang)
            except Exception as e:
                print(f"  FAILED {sheet!r} lesson {lid} ({e.__class__.__name__}: {e})")
                continue

            for r, text in zip(rows, translated):
                ws.cell(row=r, column=lang_col, value=text)

            wb.save(DB_PATH)
            done += 1
            print(f"  {sheet!r} lesson {lid} ({len(rows)} phrases) done")

            if args.sleep:
                time.sleep(args.sleep)
            if args.limit and done >= args.limit:
                print(f"Hit --limit {args.limit}, stopping.")
                return

    print(f"Done. Translated {done} lessons -> {DB_PATH}")


if __name__ == "__main__":
    main()
