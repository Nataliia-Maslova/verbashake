"""
scripts/generate_grammar_language.py -- one-time (resumable) translation of
the Grammar module's phrase database (data/imlls_database_with_titles.xlsx,
"phrases" sheet) into a brand-new target language column.

Context (2026-09-17): every language currently supported by Grammar has its
own fixed column in this workbook (en, uk, es, ko, fr, de, ja, zh, pt, it,
pl, ru, ca, nl) -- engine.loader.load_phrases() requires BOTH the native and
target columns to exist and be filled for a lesson to show up at all. Adding
a new language (e.g. Romanian) to engine.loader.LANG_COLUMNS without also
populating a real column for it would either KeyError (column missing
entirely) or silently show 0 lessons (column present but empty) -- this
script does the actual translation work that makes the new language usable
in Grammar.

Batches by lesson_id (not phrase_id) -- one Gemini call translates every
phrase in a lesson (~8) PLUS the lesson's topic label together, in a single
JSON response. This is ~182 calls total instead of ~1450+182, and keeps the
translations for one lesson mutually consistent (same register/vocabulary
across the lesson's 8 phrases), the same rationale scripts/generate_vocab_
from_cefrj.py used for picking one construction per word rather than
translating in total isolation.

Usage:
    python scripts/generate_grammar_language.py --lang Romanian --limit 3
        Small test batch -- sanity-check quality on 3 lessons first.

    python scripts/generate_grammar_language.py --lang Romanian
        Full run, all lessons missing the "ro" column. Resumable: skips any
        lesson_id whose target-language cells are already fully populated,
        so an interrupted run can just be restarted. Pass --overwrite to
        force-regenerate everything instead.

Requires GEMINI_API_KEY (real, one-time cost -- ~182 calls for a full run,
not recurring). Makes a timestamped backup of the workbook before writing
(data/imlls_database_with_titles.xlsx.bak_<YYYYMMDD>), once per day, the
same convention already used elsewhere in this project (e.g. the 2026-08-23
gender-agreement fix).
"""
from __future__ import annotations

import argparse
import datetime
import shutil
import sys
import time
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from engine import gemini
from engine.loader import LANG_COLUMNS

DB_PATH = ROOT / "data" / "imlls_database_with_titles.xlsx"


def _backup_once() -> None:
    stamp = datetime.date.today().strftime("%Y%m%d")
    bak = DB_PATH.with_name(f"{DB_PATH.name}.bak_{stamp}")
    if not bak.exists():
        shutil.copy2(DB_PATH, bak)
        print(f"Backed up {DB_PATH.name} -> {bak.name}")


def _col_letter(idx: int) -> str:
    """1-based column index -> Excel letter (A, B, ..., AA, ...)."""
    letters = ""
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _ensure_column(ws, headers: list, name: str) -> int:
    """Return the 1-based column index for `name`, appending a new header
    column at the end of the sheet if it doesn't already exist."""
    if name in headers:
        return headers.index(name) + 1
    col = len(headers) + 1
    ws.cell(row=1, column=col, value=name)
    headers.append(name)
    return col


def translate_lesson(phrases: list[str], topic_en: str, lang: str) -> dict:
    numbered = "\n".join(f"{i+1}. {p}" for i, p in enumerate(phrases))
    prompt = (
        f"Translate this English grammar-lesson topic label and its "
        f"numbered example sentences into {lang}, for a language-learning "
        f"app. Keep the same register (neutral, everyday), keep any names "
        f"(e.g. \"Alex\") as-is, and keep the same grammatical point the "
        f"English illustrates wherever {lang} has an equivalent construction. "
        f"Return ONLY a JSON object, no markdown fences, no commentary:\n"
        f'{{"topic": "<translated topic label>", '
        f'"phrases": ["<translation 1>", "<translation 2>", ...]}}\n\n'
        f"Topic: {topic_en}\n\nSentences:\n{numbered}"
    )
    result = gemini._model(gemini._LITE).generate_content(prompt)
    parsed = gemini._parse_json(gemini._safe_text(result), fallback=None)
    if not parsed or "phrases" not in parsed or len(parsed["phrases"]) != len(phrases):
        raise ValueError(f"bad/mismatched response: {parsed!r}")
    return parsed


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--lang", required=True, help="Full language name, e.g. 'Romanian'.")
    parser.add_argument("--limit", type=int, default=None, help="Max lessons to translate this run.")
    parser.add_argument("--overwrite", action="store_true", help="Regenerate even already-filled lessons.")
    parser.add_argument("--sleep", type=float, default=0.0, help="Seconds to sleep between calls.")
    args = parser.parse_args()

    if args.lang not in LANG_COLUMNS:
        sys.exit(f"'{args.lang}' not in engine.loader.LANG_COLUMNS -- add it there first.")
    code = LANG_COLUMNS[args.lang]

    _backup_once()

    wb = load_workbook(DB_PATH)
    ws = wb["phrases"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    en_col = headers.index("en") + 1
    topic_en_col = headers.index("topic_en") + 1
    lesson_id_col = headers.index("lesson_id") + 1
    lang_col = _ensure_column(ws, headers, code)
    topic_col = _ensure_column(ws, headers, f"topic_{code}")

    # Group data rows by lesson_id.
    lessons: dict[int, list[int]] = {}
    for row in range(2, ws.max_row + 1):
        lid = ws.cell(row=row, column=lesson_id_col).value
        if lid is None:
            continue
        lessons.setdefault(int(lid), []).append(row)

    done = 0
    for lid in sorted(lessons):
        rows = lessons[lid]
        already_filled = all(
            (ws.cell(row=r, column=lang_col).value or "").strip() not in ("", "TODO")
            for r in rows
        )
        if already_filled and not args.overwrite:
            continue

        en_phrases = [ws.cell(row=r, column=en_col).value or "" for r in rows]
        topic_en = ws.cell(row=rows[0], column=topic_en_col).value or ""

        try:
            translated = translate_lesson(en_phrases, topic_en, args.lang)
        except Exception as e:
            print(f"  FAILED lesson {lid} ({e.__class__.__name__}: {e})")
            continue

        for r, text in zip(rows, translated["phrases"]):
            ws.cell(row=r, column=lang_col, value=text)
            ws.cell(row=r, column=topic_col, value=translated["topic"])

        wb.save(DB_PATH)
        done += 1
        print(f"  lesson {lid} ({len(rows)} phrases) -> {translated['topic']!r}")

        if args.sleep:
            time.sleep(args.sleep)
        if args.limit and done >= args.limit:
            print(f"Hit --limit {args.limit}, stopping.")
            return

    print(f"Done. Translated {done} lessons -> {DB_PATH}")


if __name__ == "__main__":
    main()
